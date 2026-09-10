"""Create Brackett-themed benchmark charts from run-level Excel or CSV data."""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
import json
import math
from pathlib import Path
import re
from statistics import mean
import textwrap


# Brackett apps/ui/base-ui/src/tokens/tokens.css (2026-09-10).
AMBER = ["#f5b840", "#d69a2e", "#a87421", "#6b4915"]
THEMES = {
    "dark": dict(background="#0a131f", foreground="#f7f5f0", muted="#8499b0",
                 card="#131e2c", border="#243550"),
    "light": dict(background="#ffffff", foreground="#0d1824", muted="#5c7389",
                  card="#f7f5f0", border="#d8dde2"),
}


@dataclass(frozen=True)
class Run:
    case: str
    agent: str
    score: float
    completed: bool


@dataclass
class Dataset:
    runs: list[Run]
    cases: list[str]
    agents: list[str]


def score_value(value: object) -> tuple[float, bool] | None:
    """Blank is untested; NA is an attempted, incomplete run worth zero."""
    if value is None or str(value).strip() == "":
        return None
    token = str(value).strip()
    if token.casefold() in {"na", "n/a"}:
        return 0.0, False
    if isinstance(value, bool):
        raise ValueError("Boolean values are not scores")
    try:
        number = float(token[:-1]) / 100 if token.endswith("%") else float(token)
    except ValueError as exc:
        raise ValueError(f"Invalid score {value!r}; use 0–1, a percentage, NA, or blank") from exc
    if not math.isfinite(number) or not 0 <= number <= 1:
        raise ValueError(f"Score {value!r} is outside 0–1; write 86% rather than 86")
    return number, True


def parse_rows(table: list[list[object]]) -> Dataset:
    table = [row for row in table if any(v is not None and str(v).strip() for v in row)]
    if not table:
        raise ValueError("Input contains no data")
    headers = [str(v or "").strip() for v in table[0]]
    normalized = [re.sub(r"[\s_-]", "", h).casefold() for h in headers]
    case_col = next((i for i, h in enumerate(normalized) if h == "usecase"), None)
    if case_col is None:
        raise ValueError("Missing Usecase column")
    long_format = "agent" in normalized and "score" in normalized
    groups: list[tuple[str, list[int]]] = []
    if long_format:
        agent_col, score_col = normalized.index("agent"), normalized.index("score")
        run_col = next((i for i, h in enumerate(normalized) if h in {"run", "runnumber", "runid"}), None)
    else:
        for i, header in enumerate(headers):
            match = re.fullmatch(r"Avg\.?\s+(.+?)\s+Score", header, re.I)
            if match:
                indices = []
                for j in range(i + 1, len(headers)):
                    if re.fullmatch(r"Run\s+\d+", headers[j], re.I):
                        indices.append(j)
                    else:
                        break
                if not indices:
                    raise ValueError(f"{header}: run columns are required to count incomplete attempts")
                groups.append((match[1].strip(), indices))
        if not groups:
            raise ValueError("Expected Usecase/Agent/Score columns or Avg. <agent> Score followed by Run columns")
        if len({name.casefold() for name, _ in groups}) != len(groups):
            raise ValueError("Duplicate agent groups")

    runs, cases, agents = [], [], []
    seen_cases, seen_runs = set(), set()
    for line_number, row in enumerate(table[1:], 2):
        row = row + [None] * max(0, len(headers) - len(row))
        case = str(row[case_col] or "").strip()
        if not long_format and case.casefold() == "final score":
            break
        if not case:
            raise ValueError(f"Row {line_number}: missing Usecase")
        if case not in cases:
            cases.append(case)
        if long_format:
            agent = str(row[agent_col] or "").strip()
            if not agent:
                raise ValueError(f"Row {line_number}: missing Agent")
            row_groups = [(agent, [score_col])]
            if run_col is not None and row[run_col] is not None and str(row[run_col]).strip():
                key = (case, agent, str(row[run_col]).strip())
                if key in seen_runs:
                    raise ValueError(f"Row {line_number}: duplicate run {key}")
                seen_runs.add(key)
        else:
            if case in seen_cases:
                raise ValueError(f"Row {line_number}: duplicate use case {case!r}")
            seen_cases.add(case)
            row_groups = groups
        for agent, indices in row_groups:
            if agent not in agents:
                agents.append(agent)
            for i in indices:
                try:
                    parsed = score_value(row[i])
                except ValueError as exc:
                    raise ValueError(f"Row {line_number}, {headers[i]}: {exc}") from exc
                if parsed is not None:
                    runs.append(Run(case, agent, *parsed))
    if not runs:
        raise ValueError("No attempted runs found")
    preferred = ["Brackett", "Claude", "Codex"]
    agents.sort(key=lambda a: (preferred.index(a) if a in preferred else len(preferred), a))
    return Dataset(runs, cases, agents)


def load_dataset(path: Path, sheet: str = "Overview") -> Dataset:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as stream:
            return parse_rows(list(csv.reader(stream)))
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet not in workbook.sheetnames:
                raise ValueError(f"Sheet {sheet!r} not found; available: {', '.join(workbook.sheetnames)}")
            return parse_rows([list(row) for row in workbook[sheet].iter_rows(values_only=True)])
        finally:
            workbook.close()
    raise ValueError("Input must be .xlsx or .csv")


def summarize(data: Dataset) -> dict:
    scores = {a: {} for a in data.agents}
    for run in data.runs:
        scores[run.agent].setdefault(run.case, []).append(run.score)
    case_means = {a: {c: mean(v) for c, v in cases.items()} for a, cases in scores.items()}
    shared = [c for c in data.cases if all(c in case_means[a] for a in data.agents)]
    tested = [c for c in data.cases if any(c in case_means[a] for a in data.agents)]
    agents = {}
    for agent in data.agents:
        runs = [r for r in data.runs if r.agent == agent]
        agents[agent] = {
            "average_score": mean(case_means[agent].values()) if runs else None,
            "completed_runs": sum(r.completed for r in runs),
            "attempted_runs": len(runs),
            "completion_rate": sum(r.completed for r in runs) / len(runs) if runs else None,
            "tested_use_cases": len(case_means[agent]),
            "shared_score": mean(case_means[agent][c] for c in shared) if shared else None,
        }
    return dict(agents=agents, case_scores=case_means, shared_cases=shared,
                tested_cases=tested, untested_cases=[c for c in data.cases if c not in tested])


def render(data: Dataset, output: Path, source: str, theme: str = "dark",
           font: Path | None = None, cases_per_page: int = 10,
           title: str = "ShowTell benchmark") -> list[Path]:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.ticker import PercentFormatter

    family = "DejaVu Sans"
    if font:
        font_manager.fontManager.addfont(str(font))
        family = font_manager.FontProperties(fname=str(font)).get_name()
    elif any(f.name == "Satoshi" for f in font_manager.fontManager.ttflist):
        family = "Satoshi"
    palette = THEMES[theme]
    plt.rcParams.update({"font.family": family, "font.size": 11,
                         "text.color": palette["foreground"], "svg.fonttype": "path"})
    summary = summarize(data)
    output.mkdir(parents=True, exist_ok=True)
    files = []
    colors = [AMBER[i % len(AMBER)] for i in range(len(data.agents))]

    def save(fig, name):
        for extension in ("png", "svg"):
            path = output / f"{name}.{extension}"
            fig.savefig(path, dpi=180, facecolor=palette["background"])
            files.append(path)
        plt.close(fig)

    def axes_style(ax, horizontal=False):
        ax.set_facecolor(palette["card"])
        for spine in ax.spines.values():
            spine.set_visible(False)
        ax.tick_params(colors=palette["muted"], length=0, pad=10)
        axis = ax.xaxis if horizontal else ax.yaxis
        axis.set_major_formatter(PercentFormatter(1, decimals=0))
        axis.set_ticks([0, .25, .5, .75, 1])
        ax.grid(axis="x" if horizontal else "y", color=palette["border"], linewidth=.65)
        ax.set_axisbelow(True)

    fig, axes = plt.subplots(1, 3, figsize=(18, 10), facecolor=palette["background"])
    fig.subplots_adjust(left=.065, right=.96, top=.66, bottom=.23, wspace=.28)
    fig.text(.065, .92, "brackett", fontsize=22, color=AMBER[0], weight="medium")
    fig.text(.065, .845, f"{len(data.agents)} agents, compared", fontsize=38, weight="medium")
    fig.text(.065, .793, title, fontsize=15, color=palette["muted"])
    shared_count = len(summary["shared_cases"])
    panels = [("Average score", "Incomplete runs included as zero", "average_score"),
              ("Completion rate", "Completed runs / attempted runs", "completion_rate"),
              ("Shared-use-case score", f"Same {shared_count} use cases for every agent", "shared_score")]
    for ax, (heading, subtitle, key) in zip(axes, panels):
        axes_style(ax)
        ax.set_ylim(0, 1.14)
        ax.set_xlim(-.65, len(data.agents) - .35)
        ax.text(0, 1.23, heading, transform=ax.transAxes, fontsize=21, weight="medium")
        ax.text(0, 1.15, subtitle, transform=ax.transAxes, fontsize=11, color=palette["muted"])
        ax.text(0, 1.07, "HIGHER IS BETTER", transform=ax.transAxes, fontsize=10, color=palette["muted"])
        labels = []
        for i, (agent, color) in enumerate(zip(data.agents, colors)):
            record = summary["agents"][agent]
            value = record[key]
            if value is not None:
                ax.bar(i, value, width=.48, color=color, zorder=3)
                ax.text(i, value+.035, f"{value:.1%}", ha="center", fontsize=21, weight="medium")
            else:
                ax.text(i, .05, "No data", ha="center", color=palette["muted"])
            count = (f"{record['completed_runs']} / {record['attempted_runs']} runs" if key == "completion_rate"
                     else f"{shared_count if key == 'shared_score' else record['tested_use_cases']} cases")
            labels.append(f"{agent}\n{count}")
        ax.set_xticks(range(len(labels)), labels, fontsize=12)
    fig.text(.065, .125, "NA = incomplete run, scored as zero. Blank runs are untested and excluded.", fontsize=12, color=palette["muted"])
    fig.text(.065, .089, "Scores average runs within each use case, then use cases equally. Numeric scores count as completed runs.", fontsize=11, color=palette["muted"])
    fig.text(.065, .05, f"Source: {source}   ·   {len(summary['untested_cases'])} entirely untested cases omitted", fontsize=10, color=palette["muted"])
    save(fig, "benchmark-overview")

    tested = summary["tested_cases"]
    pages = math.ceil(len(tested) / cases_per_page)
    for page in range(pages):
        cases = tested[page*cases_per_page:(page+1)*cases_per_page]
        fig, ax = plt.subplots(figsize=(18, max(6, len(cases)*.95+3)), facecolor=palette["background"])
        fig.subplots_adjust(left=.36, right=.91, top=.81, bottom=.14)
        axes_style(ax, horizontal=True)
        ax.set_xlim(0, 1.13)
        ax.set_ylim(len(cases)-.5, -.7)
        ax.xaxis.tick_top()
        step = .75 / len(data.agents)
        for i, (agent, color) in enumerate(zip(data.agents, colors)):
            for j, case in enumerate(cases):
                y = j + (i-(len(data.agents)-1)/2)*step
                value = summary["case_scores"][agent].get(case)
                if value is not None:
                    ax.barh(y, value, height=step*.75, color=color, label=agent if j == 0 else None)
                    ax.text(value+.01, y, f"{value:.1%}", va="center", fontsize=10)
                else:
                    ax.text(.01, y, "Untested", va="center", fontsize=9, color=palette["muted"])
        ax.set_yticks(range(len(cases)), [textwrap.fill(c, 40) for c in cases], fontsize=12)
        fig.text(.06, .945, "brackett", fontsize=19, color=AMBER[0], weight="medium")
        fig.text(.06, .89, f"Scores by use case   {page+1} / {pages}", fontsize=28, weight="medium")
        from matplotlib.patches import Patch
        fig.legend(handles=[Patch(color=c, label=a) for a, c in zip(data.agents, colors)],
                   loc="upper right", bbox_to_anchor=(.93, .957), ncol=len(data.agents), frameon=False)
        fig.text(.06, .075, "Average run score per use case. NA runs count as zero; blank runs are excluded. Higher is better.", fontsize=12, color=palette["muted"])
        fig.text(.06, .04, f"Source: {source}", fontsize=10, color=palette["muted"])
        save(fig, f"use-case-scores-{page+1}")
    summary.update(source=source, theme=theme, font=family,
                   methodology="NA=zero/incomplete; blank=untested; equal-weight use-case means; numeric scores=completed")
    (output / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    return files


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help=".xlsx or .csv containing run-level scores")
    parser.add_argument("--sheet", default="Overview", help="Excel worksheet (default: Overview)")
    parser.add_argument("--output-dir", type=Path, default=Path("runs/benchmark-charts"))
    parser.add_argument("--theme", choices=THEMES, default="dark")
    parser.add_argument("--font", type=Path, help="Optional Satoshi .ttf/.otf file; otherwise use installed Satoshi or DejaVu Sans")
    parser.add_argument("--cases-per-page", type=int, default=10)
    parser.add_argument("--title", default="ShowTell benchmark")
    args = parser.parse_args(argv)
    if not 1 <= args.cases_per_page <= 20:
        parser.error("--cases-per-page must be between 1 and 20")
    try:
        data = load_dataset(args.input, args.sheet)
        source = args.input.name + (f" · {args.sheet}" if args.input.suffix.lower() == ".xlsx" else "")
        files = render(data, args.output_dir, source, args.theme, args.font, args.cases_per_page, args.title)
    except (ValueError, OSError, ImportError) as exc:
        parser.exit(2, f"Error: {exc}\nInstall chart dependencies with: pip install -e '.[charts]'\n")
    print(f"Created {len(files)} charts in {args.output_dir.resolve()} (PNG + SVG), plus summary.json")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
