"""Reproducible research figures from the benchmark workbook or exported attempt CSV."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from hashlib import sha256
import json
from pathlib import Path
import re

from .publication import publication_comment

from .benchmark_charts import load_dataset, score_value, summarize, THEMES, AMBER, COMPANION_COLORS


def task_area(case):
    """Editorial grouping of workbook case names; not an inferred difficulty rating."""
    name = re.sub(r' 1$', '', case)
    groups = {
        'Hiring & people': ['Candidate–Role Rematching','Job Offer Follow-Up','Job Requisition Broadcast','Job Requisition Triage','Recruiter Performance Scorecard','Leave Request Eligibility Review','New-Hire Onboarding Coordination'],
        'Inventory & orders': ['Limited-Stock Order Allocation','Low-Stock Replenishment Sweep','Order Request Replenishment','Shortage Substitution Queue'],
        'Finance & procurement': ['Credit Release Queue','Accounts Receivable Collections Follow-Up','Email Purchase Invoice Posting','Purchase Invoice Reconciliation','RFQ Quote Award','Supplier Price Increase Review','Goods Receipt Shortfall Follow-Up'],
        'Customer & sales': ['Customer Returns Inbox Triage','Discount Request Review','Customer Support Issue Triage','Inbound Sales Enquiry Routing','Return Eligibility and RMA','Returns Inbox Sweep','Warranty Claim Intake'],
        'Logistics & fleet': ['Freight Consolidation Planning','Vehicle Maintenance Service Booking','Weather-Aware Departure Review'],
    }
    return next((group for group, names in groups.items() if name in names), 'Unclassified')


EXCLUDED_RUN = re.compile(r'earlier\s+\d+', re.I)


def publication_model(agent, value):
    """Use the author's public Brackett description; preserve other source labels."""
    if agent == 'Brackett':
        return 'Fused/multiple models'
    return str(value or 'Not recorded').strip() or 'Not recorded'


def model_labels(records, agents):
    """Publication labels for selected attempts, never inferred from grader notes."""
    return {agent: dict(Counter(r['model'] for r in records if r['agent'] == agent))
            for agent in agents}


def model_note(labels):
    parts = []
    for agent, counts in labels.items():
        entries = ', '.join(f'{label} ({count} selected attempts)' for label, count in counts.items())
        parts.append(f'{agent}: {entries}.')
    return 'System/model labels: ' + ' '.join(parts) + ' Brackett uses the author-specified fused/multiple-model description. Other model labels are workbook-reported, not verified model IDs; model and adapter versions are not fully pinned. Grader names in comments do not identify the evaluated model.'


def load_attempts(path: Path):
    """Return (dataset, attempt records, excluded detail rows).

    Detailed Analysis rows are joined to Overview runs by exact system, case and
    integer run number. Rows labelled "Earlier N" are attempts the workbook
    explicitly excludes from the overview; they are returned separately for
    provenance and never enter the scores.
    """
    data = load_dataset(path)
    details, excluded = {}, []
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        if 'Detailed Analysis' in book.sheetnames:
            rows = list(book['Detailed Analysis'].values)
            headers = [str(v or '').strip() for v in rows[0]]
            for row in rows[1:]:
                item = dict(zip(headers, row))
                if not item.get('Agent'): continue
                label = str(item['Run Number']).strip()
                note = str(item.get('Comments') or '')
                model = publication_model(item['Agent'], item.get('Model'))
                if EXCLUDED_RUN.fullmatch(label):
                    if 'excluded from overview' not in note.lower():
                        raise ValueError(f'Run {label!r} for {item["Agent"]}/{item["Usecase"]} is not an overview run; the note must say it is excluded from overview')
                    score = score_value(item['Score'])
                    excluded.append(dict(agent=item['Agent'], case=item['Usecase'], label=label,
                                         score=None if score is None else score[0], completed=bool(score and score[1]),
                                         model=model, comment=publication_comment(note)))
                    continue
                try:
                    number = int(float(label))
                except ValueError as exc:
                    raise ValueError(f'Unrecognised run label {label!r} for {item["Agent"]}/{item["Usecase"]}') from exc
                key = (item['Agent'], item['Usecase'], number)
                if key in details: raise ValueError(f'Duplicate detail: {key}')
                details[key] = (score_value(item['Score']), note, model)
        book.close()
    elif path.suffix.lower() == '.csv':
        with path.open(newline='') as stream:
            for item in csv.DictReader(stream):
                if 'Comment' in item or 'Model' in item:
                    key = (item['Agent'], item['Usecase'], int(item['Run']))
                    details[key] = (score_value(item['Score']), item.get('Comment', ''), publication_model(item['Agent'], item.get('Model')))
    records, counts = [], Counter()
    for run in data.runs:
        counts[run.agent, run.case] += 1
        number = counts[run.agent, run.case]
        key = (run.agent, run.case, number)
        note = ''
        model = publication_model(run.agent, None)
        if key in details:
            score, note, model = details[key]
            # Excel may store 0.453 as 0.45299999999999996; tolerate representation noise only.
            if score is None or score[1] != run.completed or abs(score[0] - run.score) > 1e-9:
                raise ValueError(f'Overview and detail disagree: {key}')
            note = publication_comment(note)
        if run.completed: outcome = 'Scored'
        elif note.lower().startswith('empty shortcut'): outcome = 'Empty shortcut'
        elif note.lower().startswith('refusal'): outcome = 'Refusal'
        else: outcome = 'Other incomplete'
        records.append(dict(agent=run.agent, case=run.case, run=number, score=run.score,
                            completed=run.completed, outcome=outcome, comment=note, model=model))
    return data, records, excluded


def analyze(data, records, task_groups=None, excluded=()):
    result = summarize(data)
    result['model_labels'] = model_labels(records, data.agents)
    result['outcomes'] = {a: dict(Counter(r['outcome'] for r in records if r['agent'] == a)) for a in data.agents}
    pairs = defaultdict(list)
    for r in records: pairs[r['agent'], r['case']].append(r)
    # Attempts per case differ by system and case; record the distribution rather than assuming a schedule.
    result['attempts_per_case'] = {a: {str(n): k for n, k in sorted(Counter(len(runs) for (agent, _), runs in pairs.items() if agent == a).items())} for a in data.agents}
    result['repeat_gaps'] = []
    for (agent, case), runs in pairs.items():
        if len(runs) < 2: continue
        scores = [r['score'] for r in runs]
        result['repeat_gaps'].append(dict(agent=agent, case=case, attempts=len(runs),
            gap=max(scores)-min(scores), scores=scores, completed=[r['completed'] for r in runs]))
    result['excluded_attempts'] = list(excluded)
    wins = Counter()
    for case in result['shared_cases']:
        best = max(result['case_scores'][a][case] for a in data.agents)
        leaders = [a for a in data.agents if abs(result['case_scores'][a][case]-best) < 1e-12]
        wins[leaders[0] if len(leaders) == 1 else 'Tie'] += 1
    result['shared_case_leaders'] = dict(wins)
    # Sorted so a workbook and its exported attempt CSV yield the same inventory order.
    result['task_inventory'] = [dict(case=c,area=(task_groups or {}).get(c, task_area(c)),tested=c in result['tested_cases']) for c in sorted(data.cases)]
    return result


def render(data, records, result, output: Path, theme='light', task_quiz=None):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import numpy as np
    from matplotlib.colors import LinearSegmentedColormap
    from matplotlib.patches import Patch
    from matplotlib.ticker import PercentFormatter
    p = THEMES[theme]
    colors = {'Brackett': AMBER[0], 'Claude': COMPANION_COLORS[0], 'Codex': COMPANION_COLORS[1]}
    agents = sorted(data.agents, key=lambda a: (['Brackett', 'Codex', 'Claude'].index(a) if a in colors else 3, a))
    plt.rcParams.update({'font.family':'DejaVu Sans', 'font.size':13, 'text.color':p['foreground'],
                         'axes.labelcolor':p['foreground'], 'xtick.color':p['muted'],
                         'ytick.color':p['foreground'], 'svg.fonttype':'none', 'svg.hashsalt':'showtell-analysis'})
    output.mkdir(parents=True, exist_ok=True)
    def save(fig, name):
        for ext in ('svg', 'png'):
            fig.savefig(output / f'{name}.{ext}', facecolor=p['background'], dpi=170,
                        metadata={'Date':None} if ext == 'svg' else {})
        # Keep SVG text selectable, with a browser-safe fallback on machines without DejaVu.
        svg_path = output / f'{name}.svg'
        svg_path.write_text(svg_path.read_text().replace("font-family: 'DejaVu Sans'", "font-family: 'DejaVu Sans', Arial, sans-serif"))
        plt.close(fig)
    def style(ax):
        ax.set_facecolor(p['background'])
        for s in ax.spines.values(): s.set_visible(False)
        ax.tick_params(length=0, pad=9)
        ax.set_axisbelow(True)
    def title(fig, headline, subtitle):
        fig.text(.04,.96,headline,fontsize=21,weight='medium',va='top')
        fig.text(.04,.915,subtitle,fontsize=12,color=p['muted'],va='top')

    # The same hue scale for all systems; missing values never become zero.
    cases = sorted(result['tested_cases'])
    values = np.array([[result['case_scores'][a].get(c, np.nan) for a in agents] for c in cases])
    fig, ax = plt.subplots(figsize=(12, max(8, .34*len(cases)+2.4)))
    fig.subplots_adjust(left=.42,right=.88,top=.855,bottom=.09)
    cmap = LinearSegmentedColormap.from_list('brackett-score', ['#fffaf0','#f5b840','#835407'])
    cmap.set_bad('#e3e7eb')
    ax.imshow(np.ma.masked_invalid(values), cmap=cmap, vmin=0, vmax=1, aspect='auto')
    ax.set_xticks(range(len(agents)), agents, fontsize=15)
    ax.xaxis.tick_top()
    ax.set_yticks(range(len(cases)), cases, fontsize=12)
    ax.tick_params(length=0,pad=10)
    for spine in ax.spines.values(): spine.set_visible(False)
    for i, case in enumerate(cases):
        for j, agent in enumerate(agents):
            score=values[i,j]
            label='—' if np.isnan(score) else f'{score*100:.1f}'
            ax.text(j,i,label,ha='center',va='center',fontsize=12,
                    color='#fff' if score > .86 else '#0d1824')
    ax.set_yticks(np.arange(-.5,len(cases),1),minor=True)
    ax.grid(which='minor',axis='y',color=p['background'],linewidth=2)
    ax.tick_params(which='minor',length=0)
    title(fig,'Every tested business process',f'{len(cases)} cases · mean comprehension score (%) · identical color scale for every system')
    cax=fig.add_axes([.915,.2,.016,.5])
    fig.colorbar(ax.images[0], cax=cax, ticks=[0,.25,.5,.75,1],format=PercentFormatter(1))
    fig.text(.04,.04,'NA attempts count as zero. Grey / — = untested. Source case names, including “1”, are preserved.',fontsize=11,color=p['muted'])
    save(fig,'case-heatmap')

    fig, ax = plt.subplots(figsize=(11,5.3))
    fig.subplots_adjust(left=.15,right=.94,top=.72,bottom=.22)
    style(ax)
    outcomes=['Scored','Empty shortcut','Refusal','Other incomplete']
    fills={'Empty shortcut':'#d8dde2','Refusal':'#6b7785','Other incomplete':'#a7afb8'}
    for y,agent in enumerate(agents):
        counts=result['outcomes'][agent]; total=sum(counts.values()); left=0
        for outcome in outcomes:
            n=counts.get(outcome,0)
            if not n: continue
            value=n/total
            ax.barh(y,value,left=left,height=.46,color=colors.get(agent,AMBER[0]) if outcome=='Scored' else fills[outcome])
            ax.text(left+value/2,y,str(n),ha='center',va='center',fontsize=15,
                    color='white' if outcome=='Refusal' else '#0d1824')
            left+=value
    ax.set_yticks(range(len(agents)),[f'{a}\nn={sum(result["outcomes"][a].values())}' for a in agents]);ax.invert_yaxis()
    ax.set_xlim(0,1);ax.xaxis.set_major_formatter(PercentFormatter(1));ax.set_xticks([0,.25,.5,.75,1]);ax.grid(axis='x',color=p['border'])
    title(fig,'Was there an answer to grade?','All tested cases · bars show proportions; labels show attempt counts')
    fig.legend(handles=[Patch(color='#f5b840',label='Scored (system color)')]+[Patch(color=fills[o],label=o) for o in outcomes[1:] if any(result['outcomes'][a].get(o) for a in agents)],loc='lower center',bbox_to_anchor=(.5,.075),ncol=3,frameon=False,fontsize=11)
    fig.text(.04,.025,'“Scored” is not task success. Incomplete reasons come from workbook notes; they are not independently adjudicated.',fontsize=10,color=p['muted'])
    save(fig,'attempt-outcomes')

    # Highlight the largest observed spreads, not a learning trend or confidence interval.
    repeat_agents=[a for a in agents if any(x['agent']==a for x in result['repeat_gaps'])]
    fig, axes=plt.subplots(1,max(1,len(repeat_agents)),figsize=(13,5.8),squeeze=False)
    fig.subplots_adjust(left=.04,right=.92,top=.71,bottom=.2,wspace=.6)
    for ax,agent in zip(axes[0],repeat_agents):
        style(ax)
        rows=sorted((r for r in result['repeat_gaps'] if r['agent']==agent),key=lambda r:(-r['gap'],r['case']))[:3]
        ax.set_xlim(-.02,1.17);ax.set_ylim(len(rows)-.5,-.8)
        for y,r in enumerate(rows):
            ax.plot([min(r['scores']),max(r['scores'])],[y,y],color=colors.get(agent,AMBER[0]),linewidth=3,zorder=1)
            for i,(v,complete) in enumerate(zip(r['scores'],r['completed'])):
                if complete:
                    ax.scatter(v,y,s=90,marker='o',facecolors=colors.get(agent,AMBER[0]) if i==0 else 'white',edgecolors=colors.get(agent,AMBER[0]),linewidths=2,zorder=2)
                else:
                    ax.scatter(v,y,s=90,marker='x',color=colors.get(agent,AMBER[0]),linewidths=2,zorder=2)
            ax.text(0,y-.27,f"{r['case']}  ·  {r['attempts']} attempts",fontsize=10,color=p['foreground'])
            ax.text(1.03,y,f'  {r["gap"]*100:.1f} pp',fontsize=11,va='center')
        ax.set_yticks([]);ax.set_xticks([0,.5,1]);ax.xaxis.set_major_formatter(PercentFormatter(1));ax.grid(axis='x',color=p['border'])
        ax.set_title(agent,loc='left',fontsize=16,pad=12)
    title(fig,'The average can hide a large spread','Three largest ranges between selected attempts of one case, per system · incomplete attempts included as zero')
    fig.text(.04,.08,'● First selected attempt   ○ Later selected attempts   × Incomplete attempt',fontsize=12,color=p['muted'])
    unrepeated=[a for a in agents if a not in repeat_agents]
    plotted=sorted({r['attempts'] for r in result['repeat_gaps']})
    counts_text=' or '.join(str(n) for n in plotted)+' selected attempts per plotted case; the range is max − min.'
    note='No repeat estimate: '+', '.join(unrepeated)+'. '+counts_text if unrepeated else counts_text
    fig.text(.04,.03,'Repeated attempts, not before/after learning tests. '+note,fontsize=11,color=p['muted'])
    save(fig,'repeat-spread')

    groups=Counter(x['area'] for x in result['task_inventory'])
    groups=sorted(groups,key=lambda g:(-groups[g],g))
    fig, ax=plt.subplots(figsize=(11,5.3))
    fig.subplots_adjust(left=.29,right=.92,top=.75,bottom=.2)
    style(ax)
    for i,g in enumerate(groups):
        count=sum(x['area']==g for x in result['task_inventory'])
        ax.barh(i,count,color=AMBER[0],height=.48)
        ax.text(count+.15,i,str(count),va='center',fontsize=15)
    ax.set_yticks(range(len(groups)),groups);ax.invert_yaxis()
    ax.set_xlim(0,max(Counter(x['area'] for x in result['task_inventory']).values())+1.5)
    from matplotlib.ticker import MaxNLocator
    ax.xaxis.set_major_locator(MaxNLocator(integer=True));ax.grid(axis='x',color=p['border'])
    title(fig,'What kinds of work are in the task inventory?',f'{len(data.cases)} listed cases · grouped by business process, not model score')
    fig.text(.04,.055,f'Grouped from workbook case names; variants remain separate. {len(result["tested_cases"])} cases have attempts; {len(result["untested_cases"])} are untested.',fontsize=11,color=p['muted'])
    save(fig,'task-inventory')

    if task_quiz:
        quiz=json.loads(task_quiz.read_text())
        questions=quiz['questions']
        types=['voice_recording','screenshot','source_file']
        labels={'voice_recording':'Narration','screenshot':'Screen capture','source_file':'Source email / file'}
        names={'q_returns_inbox_cadence':'Morning cadence','q_morning_pass_scope':'Which messages belong','q_refund_delivery_window':'Delivery window','q_direct_refund_amount_limit':'Approval authority','q_direct_refund_reply_terms':'What the reply must say','q_over_limit_routing':'Where to escalate','q_finance_summary_fields':'Context for finance','q_exact_boundary_cases':'Exact boundaries','q_non_customer_inbox_items':'What to leave untouched','q_damage_override_transfer':'Damage takes precedence','q_safe_automation_error_diagnosis':'Diagnose a wrong automation'}
        matrix=np.array([[any(e['type']==t for e in q.get('evidence',[])) for t in types] for q in questions])
        fig,ax=plt.subplots(figsize=(11,max(6,.33*len(questions)+2.1)))
        fig.subplots_adjust(left=.39,right=.95,top=.76,bottom=.15)
        ax.imshow(matrix,cmap=LinearSegmentedColormap.from_list('evidence',['#f1f3f5',AMBER[0]]),vmin=0,vmax=1,aspect='auto')
        ax.set_yticks(range(len(questions)),[names.get(q['id'],q['id']) for q in questions],fontsize=12)
        ax.set_xticks(range(3),[labels[t] for t in types],fontsize=12);ax.xaxis.tick_top();ax.tick_params(length=0,pad=12)
        for spine in ax.spines.values():spine.set_visible(False)
        for i in range(len(questions)):
            for j in range(3):ax.text(j,i,'●' if matrix[i,j] else '—',ha='center',va='center',fontsize=15,color=p['foreground'] if matrix[i,j] else '#b5bdc6')
        ax.set_yticks(np.arange(-.5,len(questions),1),minor=True);ax.grid(which='minor',axis='y',color='white',linewidth=2);ax.tick_params(which='minor',length=0)
        counts=Counter(q['type'] for q in questions)
        title(fig,'Every question has a source',f'Published returns task · {len(questions)} questions · {counts["closed"]} closed answer / {counts["llm_judge"]} rubric scored')
        fig.text(.04,.055,'● Evidence cited by the question. Multiple captures of the same type count once. This is evidence coverage, not a score.',fontsize=11,color=p['muted'])
        result['question_evidence']={'source_sha256':sha256(task_quiz.read_bytes()).hexdigest(),'types':types,'ids':[q['id'] for q in questions],'matrix':matrix.astype(int).tolist(),'question_types':dict(counts)}
        save(fig,'task-evidence')


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input',type=Path)
    parser.add_argument('--output-dir',type=Path,default=Path('runs/benchmark-analysis'))
    parser.add_argument('--theme',choices=['light'],default='light',help='Light publication theme')
    parser.add_argument('--task-quiz',type=Path,help='Optional questions.json to graph task evidence coverage')
    parser.add_argument('--task-groups',type=Path,help='JSON object mapping exact case names to business-process groups')
    args=parser.parse_args(argv)
    data,records,excluded=load_attempts(args.input)
    result=analyze(data,records,json.loads(args.task_groups.read_text()) if args.task_groups else None,excluded)
    result['source']={'file':args.input.name,'sha256':sha256(args.input.read_bytes()).hexdigest(),
                      'status':'Selected attempts; non-canonical grading (manual grades and imported model-assisted grades, as recorded in the workbook notes); no uncertainty interval claimed.'}
    render(data,records,result,args.output_dir,args.theme,args.task_quiz)
    (args.output_dir/'analysis.json').write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n')
    with (args.output_dir/'attempts.csv').open('w',newline='') as stream:
        writer=csv.writer(stream,lineterminator='\n');writer.writerow(['Usecase','Agent','Run','Score','Outcome','Comment','Model'])
        for r in records:writer.writerow([r['case'],r['agent'],r['run'],r['score'] if r['completed'] else 'NA',r['outcome'],r['comment'],r['model']])
        for case in result['untested_cases']:writer.writerow([case,data.agents[0],1,'','Untested','',''])
    print(json.dumps({'outcomes':result['outcomes'],'shared_case_leaders':result['shared_case_leaders']},indent=2))
    return 0


if __name__=='__main__':raise SystemExit(main())
